import csv
import fcntl
import io
import json
import logging
import os
import re
import threading
from datetime import date, datetime
from urllib.parse import quote
from zoneinfo import ZoneInfo

from flask import Flask, Response, abort, g, jsonify, redirect, render_template, request
from date_ranges import install_date_ranges

from coverage_query import (
    allowed_store_names,
    build_coverage_view,
    collect_store_names,
    coverage_back_href,
    coverage_store_href,
    load_directory,
    store_coverage_detail,
)
from group_keyword import (
    GroupListError,
    KeywordError,
    open_ro,
    query_group_list,
    query_keyword_groups,
)
from group_query import group_events, group_info, local_context
from scrm_client import ScrmError
from scrm_sync import sync_from_scrm
from score_web import register_score
import timeout_algo
from timeout_query import query_chart, query_meta
import stc_live
from timeout_store import connect, replace_context

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

ROOT = os.path.dirname(os.path.abspath(__file__))
app = Flask(__name__)

# --- 设计系统（antd-lite.css / sprite.svg，由 design-system/sync.sh 同步到 ds/）---
_DS_DIR = os.path.join(ROOT, "ds")
_ds_cache: dict[str, str] = {}


def _ds(name: str) -> str:
    if name not in _ds_cache:
        with open(os.path.join(_DS_DIR, name), encoding="utf-8") as fh:
            _ds_cache[name] = fh.read()
    return _ds_cache[name]


app.jinja_env.globals["ds_css"] = lambda: _ds("antd-lite.css")
app.jinja_env.globals["ds_sprite"] = lambda: _ds("sprite.svg")
app.jinja_env.globals["ds_cards"] = lambda: _ds("cards.js")
app.jinja_env.globals["coverage_store_href"] = coverage_store_href

# --- m-hero SSO (Casdoor): inert unless MHERO_AUTH[_CLEANER]=audit|on is set (see docs/SSO_CASDOOR.md) ---
from mhero_auth.flask_adapter import init_app as _mhero_auth_init
_mhero_auth_init(app, "cleaner", default_mode="off", public_get_paths=["/static/brand/*"], admin_paths=["/admin/*", "/api/admin/*"])
install_date_ranges(app)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024
# 开发环境可用 CLEANER_DB_PATH / CLEANER_DATA_FILE 指向真实数据副本；默认不变
DATA_FILE = os.environ.get("CLEANER_DATA_FILE") or os.path.join(ROOT, "data.json")
GROUPS_DB_PATH = os.environ.get("CLEANER_DB_PATH") or os.path.join(ROOT, "timeout.db")
LOCK_FILE = os.path.join(ROOT, ".sync.lock")
SHANGHAI = ZoneInfo("Asia/Shanghai")
CUTOFF_DATE = datetime(2026, 5, 1, tzinfo=SHANGHAI)
_SYNC_LOCK = threading.Lock()


def _acquire_sync_lock():
    if not _SYNC_LOCK.acquire(blocking=False):
        raise ScrmError("正在同步，请稍候")
    lock_fh = open(LOCK_FILE, "w")
    try:
        fcntl.flock(lock_fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        lock_fh.close()
        _SYNC_LOCK.release()
        raise ScrmError("正在同步，请稍候")
    return lock_fh


def _release_sync_lock(lock_fh):
    fcntl.flock(lock_fh, fcntl.LOCK_UN)
    lock_fh.close()
    _SYNC_LOCK.release()


def save_data(stats, summary):
    # data.json 只保留汇总统计；同步结果里的明细和员工姓名不落盘。
    summary = {
        key: value
        for key, value in (summary or {}).items()
        if key not in {"all_rows", "filter_rows", "employee_names", "employee_ids", "robot_chat_ids"}
    }
    data = {
        "updated_at": datetime.now(SHANGHAI).strftime("%Y-%m-%d %H:%M:%S"),
        "stats": stats,
        "summary": summary,
    }
    tmp = DATA_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(tmp, DATA_FILE)


def _save_group_context(result: dict) -> None:
    conn = connect(_db_path())
    try:
        replace_context(
            conn,
            robot_chat_ids=result.get("robot_chat_ids") or [],
            employee_ids=result.get("employee_ids") or [],
            group_rows=result.get("all_rows") or [],
            employee_names=result.get("employee_names") or {},
        )
        conn.commit()
    finally:
        conn.close()


def _run_group_and_timeout():
    result = sync_from_scrm(ROOT)
    stats = compute_stats(result["all_rows"], result["filter_rows"])
    summary = result["summary"]
    try:
        _save_group_context(result)
    except Exception:
        logging.exception("保存群-门店映射失败，拉群覆盖数据仍会保存")
    save_data(stats, summary)
    return summary


def run_sync():
    """定时任务：先同步客户群，再拉会话存档按规则计算超时（同一时刻 00:00 / 12:00）。"""
    lock_fh = _acquire_sync_lock()
    try:
        summary = _run_group_and_timeout()
        try:
            summary["timeout"] = timeout_algo.run(ROOT, label="cron")
        except Exception:
            logging.exception("超时计算失败，拉群覆盖数据已保存")
        return summary
    finally:
        _release_sync_lock(lock_fh)


def run_timeout_sync(**kw):
    lock_fh = _acquire_sync_lock()
    try:
        return timeout_algo.run(ROOT, **kw)
    finally:
        _release_sync_lock(lock_fh)


def _is_admin() -> bool:
    from flask import g
    s = getattr(g, "mhero_session", None)
    if s is None:
        return False
    return getattr(s, "role", "") == "admin"


def _actor() -> str:
    from flask import g
    s = getattr(g, "mhero_session", None)
    return (getattr(s, "dn", "") or getattr(s, "name", "") or "admin") if s else "admin"


@app.context_processor
def _inject_admin():
    try:
        return {"is_admin": _is_admin()}
    except Exception:
        return {"is_admin": False}


NON_KPI_TAG = '<span class="ant-tag ant-tag-warning kpi-off" title="该规则只在超时机器人中展示，不计入 KPI/评分">仅展示，不计 KPI</span>'


@app.context_processor
def _inject_kpi():
    try:
        _keys, names = timeout_algo.non_kpi_rules(ROOT)
    except Exception:
        names = set()
    return {"non_kpi_names": sorted(names)}


@app.template_filter("dept_leaf")
def _dept_leaf(value):
    parts = [part.strip() for part in str(value or "").split("/") if part.strip()]
    return parts[-1] if parts else ""


@app.template_filter("minute_time")
def _minute_time(value):
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return text
    return parsed.strftime("%Y-%m-%d %H:%M")


@app.template_filter("kpi_tag")
def _kpi_tag(rule_name):
    from markupsafe import Markup
    try:
        names = timeout_algo.non_kpi_rules(ROOT)[1]
    except Exception:
        return Markup("")
    parts = [n.strip() for n in str(rule_name or "").split(" / ")]
    return Markup(" " + NON_KPI_TAG) if any(p in names for p in parts) else Markup("")


try:
    timeout_algo.open_db(ROOT).close()  # 建表/迁移/播种规则
except Exception:
    logging.exception("初始化超时规则表失败")


def _data_path() -> str:
    return os.environ.get("CLEANER_DATA_FILE") or DATA_FILE


def _db_path() -> str:
    return os.environ.get("CLEANER_DB_PATH") or GROUPS_DB_PATH


def load_data():
    path = _data_path()
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def parse_ct(row):
    ct = row.get("create_time")
    if not ct:
        return None
    if isinstance(ct, datetime):
        if ct.tzinfo is None:
            return ct.replace(tzinfo=SHANGHAI)
        return ct
    try:
        dt = datetime.fromisoformat(str(ct))
    except ValueError:
        return None
    if dt.tzinfo is None:
        return dt.replace(tzinfo=SHANGHAI)
    return dt


def build_monthly_stats(rows):
    months = {}
    for r in rows:
        ct = parse_ct(r)
        if ct is None:
            continue
        month_key = ct.strftime("%Y-%m")
        if month_key not in months:
            months[month_key] = {"count": 0, "stores": {}}
        months[month_key]["count"] += 1
        store = r["store"]
        months[month_key]["stores"][store] = months[month_key]["stores"].get(store, 0) + 1
    return months


def compute_stats(all_rows, filter_rows):
    all_valid = [r for r in all_rows if r.get("store")]
    filter_valid = [r for r in filter_rows if r.get("store")]

    def store_counts(rows):
        stats = {}
        for r in rows:
            store = r["store"]
            stats[store] = stats.get(store, 0) + 1
        return dict(sorted(stats.items(), key=lambda x: -x[1]))

    all_store_stats = store_counts(all_valid)
    filter_store_stats = store_counts(filter_valid)
    all_after = [r for r in all_valid if parse_ct(r) and parse_ct(r) >= CUTOFF_DATE]
    filter_after = [r for r in filter_valid if parse_ct(r) and parse_ct(r) >= CUTOFF_DATE]

    all_months_raw = build_monthly_stats(all_valid)
    filter_months_raw = build_monthly_stats(filter_valid)
    all_keys = sorted(set(all_months_raw) | set(filter_months_raw))
    years = sorted({m.split("-")[0] for m in all_keys})

    monthly_data = {}
    for month in all_keys:
        all_m = all_months_raw.get(month, {})
        filter_m = filter_months_raw.get(month, {})
        monthly_data[month] = {
            "all": {
                "group_count": all_m.get("count", 0),
                "store_count": len(all_m.get("stores", {})),
                "store_stats": dict(sorted(all_m.get("stores", {}).items(), key=lambda x: -x[1])),
            },
            "filter": {
                "group_count": filter_m.get("count", 0),
                "store_count": len(filter_m.get("stores", {})),
                "store_stats": dict(sorted(filter_m.get("stores", {}).items(), key=lambda x: -x[1])),
            },
        }

    return {
        "national": {
            "all": {
                "group_count": len(all_valid),
                "store_count": len(all_store_stats),
                "store_stats": all_store_stats,
            },
            "filter": {
                "group_count": len(filter_valid),
                "store_count": len(filter_store_stats),
                "store_stats": filter_store_stats,
            },
        },
        "after_may1": {
            "all": {
                "group_count": len(all_after),
                "store_count": len(store_counts(all_after)),
                "store_stats": store_counts(all_after),
            },
            "filter": {
                "group_count": len(filter_after),
                "store_count": len(store_counts(filter_after)),
                "store_stats": store_counts(filter_after),
            },
        },
        "monthly": {
            "months": all_keys,
            "years": years,
            "data": monthly_data,
        },
    }


@app.route("/")
def index():
    data = load_data()
    coverage = None
    if data is not None:
        coverage = build_coverage_view(
            data,
            request.args.get("theater") or "",
            request.args.get("q") or "",
            _db_path(),
        )
    return render_template("index.html", has_data=data is not None, data=data, coverage=coverage)


@app.route("/coverage/store")
def coverage_store_page():
    name = (request.args.get("name") or "").strip()
    return redirect(coverage_store_href(name, request.args.get("theater"), request.args.get("q")) if name else "/", code=302)


@app.route("/store/<path:store_key>")
def store_page(store_key):
    directory = load_directory(_db_path())
    data = load_data()
    names = {row["name"] for row in directory} | collect_store_names(data)
    name = store_key if store_key in names else next(
        (row["name"] for row in directory if row.get("code") and row["code"] == store_key), store_key
    )
    theater = (request.args.get("theater") or "").strip()
    q = (request.args.get("q") or "").strip()
    filters = {key: (request.args.get(key) or "").strip() for key in ("group_from", "group_to", "date_from", "date_to")}
    filters['date_from'], filters['date_to'] = (d.isoformat() for d in g.date_range)
    for key, value in filters.items():
        if value:
            try:
                filters[key] = date.fromisoformat(value).isoformat()
            except ValueError:
                abort(400, description="日期格式应为 YYYY-MM-DD")
    detail = None
    db_path = _db_path()
    available = os.path.isfile(db_path)
    if available and name:
        conn = open_ro(db_path)
        try:
            detail = store_coverage_detail(conn, name, q=q, **filters)
        finally:
            conn.close()
    if detail is None and name in names:
        national = ((data or {}).get("stats") or {}).get("national") or {}
        group_count = int(((national.get("all") or {}).get("store_stats") or {}).get(name) or 0)
        robot_count = int(((national.get("filter") or {}).get("store_stats") or {}).get(name) or 0)
        detail = {
            "store": name, "theater": "", "group_count": group_count,
            "robot_count": robot_count, "pending_count": max(0, group_count - robot_count),
            "pending": [], "robot": [], "events": [], "excluded": [], "snapshot_only": True,
        }
    return render_template(
        "coverage_store.html",
        detail=detail,
        name=name,
        available=available,
        back_href=coverage_back_href(theater, request.args.get("back_q") or ""),
        q=q,
        filters=filters,
        reset_href=coverage_store_href(name, theater, request.args.get("back_q")),
        nav_active="coverage",
    )


def _export_allowed(data):
    if data is None:
        return None
    return allowed_store_names(
        data,
        request.args.get("theater") or "",
        request.args.get("q") or "",
        _db_path(),
    )


@app.route("/timeout")
def timeout_page():
    data = load_data()
    return render_template("timeout.html", has_data=data is not None, data=data, start=g.date_range[0].isoformat(), end=g.date_range[1].isoformat())


@app.route("/parse")
def parse_gone():
    return redirect("/", code=301)


@app.route("/api/export")
def export_csv():
    scope = request.args.get("scope", "national")
    dataset = request.args.get("dataset", "all")
    data = load_data()
    if not data:
        return jsonify({"error": "暂无数据，请先从 SCRM 同步"}), 400

    store_stats = data["stats"].get(scope, {}).get(dataset, {}).get("store_stats", {})
    allowed = _export_allowed(data)
    if allowed is not None:
        store_stats = {name: count for name, count in store_stats.items() if name in allowed}
    if not store_stats:
        return jsonify({"error": "没有可导出的数据"}), 400

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["排名", "门店名称", "群数量"])
    for idx, (store, count) in enumerate(store_stats.items(), 1):
        writer.writerow([idx, store, count])

    filename_map = {
        ("national", "all"): "全国全量各店群数量",
        ("national", "filter"): "全国已拉机器人各店群数量",
        ("after_may1", "all"): "5月1日后全量各店群数量",
        ("after_may1", "filter"): "5月1日后已拉机器人各店群数量",
    }
    filename = filename_map.get((scope, dataset), "各店群数量") + ".csv"
    return Response(
        output.getvalue().encode("utf-8-sig"),
        mimetype="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename)},
    )


@app.route("/api/export-all")
def export_all_csv():
    data = load_data()
    if not data:
        return jsonify({"error": "暂无数据，请先从 SCRM 同步"}), 400

    all_stats = data["stats"]["national"]["all"]["store_stats"]
    filter_stats = data["stats"]["national"]["filter"]["store_stats"]
    all_stores = set(all_stats) | set(filter_stats)
    allowed = _export_allowed(data)
    if allowed is not None:
        all_stores &= allowed

    rows = []
    for store in all_stores:
        all_count = all_stats.get(store, 0)
        filter_count = filter_stats.get(store, 0)
        remaining = all_count - filter_count
        rate = f"{filter_count / all_count * 100:.1f}%" if all_count else "-"
        rows.append((store, all_count, filter_count, rate, remaining))
    rows.sort(key=lambda x: x[4], reverse=True)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["序号", "店名", "全量群数", "拉机器人群数", "拉群率", "剩余待拉群数"])
    for idx, (store, all_count, filter_count, rate, remaining) in enumerate(rows, 1):
        writer.writerow([idx, store, all_count, filter_count, rate, remaining])

    return Response(
        output.getvalue().encode("utf-8-sig"),
        mimetype="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote("门店群数量总表.csv")},
    )


@app.route("/api/timeout/meta")
def timeout_meta_api():
    return jsonify(query_meta(ROOT))


@app.route("/api/timeout/chart")
def timeout_chart_api():
    payload = query_chart(
        ROOT,
        scope=request.args.get("scope") or "national",
        ident=request.args.get("id") or "",
        grain=request.args.get("grain") or "month",
        start=g.date_range[0].isoformat(),
        end=g.date_range[1].isoformat(),
        theater=request.args.get("theater") or "",
        q=request.args.get("q") or "",
    )
    return jsonify(payload)


_ROOM_RE = re.compile(r"^[A-Za-z0-9_\-]{6,80}$")


@app.route("/group/<room_id>")
def group_page(room_id):
    start, end = (day.isoformat() for day in g.date_range)
    db = _db_path()
    info = None
    groups = {"kpi_counted": [], "kpi_excluded": [], "display_only": []}
    if _ROOM_RE.match(room_id) and os.path.isfile(db):
        conn = open_ro(db)
        try:
            info = group_info(conn, room_id)
            if info is not None:
                groups = group_events(os.path.dirname(os.path.abspath(db)), conn, room_id, *g.date_range)
        finally:
            conn.close()
    if info is None:
        return (
            render_template(
                "group.html",
                info=None,
                groups=None,
                counts=None,
                latest="",
                start=start,
                end=end,
                back_href="/",
                back_label="拉群覆盖",
                nav_active="coverage",
            ),
            404,
        )
    counts = {key: len(value) for key, value in groups.items()}
    latest = max(
        (event.get("created_at") or "" for items in groups.values() for event in items),
        default="",
    )
    if info.get("store"):
        back_href = (
            "/store/" + quote(info["store"], safe="")
            + "?start=" + quote(start) + "&end=" + quote(end)
        )
        back_label = info["store"]
    else:
        back_href, back_label = "/", "拉群覆盖"
    return render_template(
        "group.html",
        info=info,
        groups=groups,
        counts=counts,
        latest=latest,
        start=start,
        end=end,
        back_href=back_href,
        back_label=back_label,
        nav_active="coverage",
    )


@app.route("/api/group/<room_id>/context")
def group_context_api(room_id):
    if not _ROOM_RE.match(room_id):
        abort(400)
    at = (request.args.get("at") or "").strip()
    if not re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$", at):
        abort(400)
    try:
        minutes = int(request.args.get("minutes") or 30)
    except ValueError:
        minutes = 30
    minutes = min(max(minutes, 1), 240)
    msg_id = (request.args.get("msg_id") or "").strip()[:120]
    local = {"messages": [], "source": "local"}
    db = _db_path()
    if os.path.isfile(db):
        conn = open_ro(db)
        try:
            robot_ids = {part.strip() for part in os.environ.get("SCRM_ROBOT_USERIDS", "").split(",") if part.strip()}
            try:
                employee_ids = [row["userid"] for row in conn.execute("SELECT userid FROM employee")]
            except Exception:
                employee_ids = []
            local = local_context(conn, room_id, msg_id, at, minutes, employee_ids, robot_ids)
        except Exception:
            logging.exception("读取本地聊天上下文失败")
            local = {"messages": [], "source": "local"}
        finally:
            conn.close()
    if local.get("messages"):
        local["source"] = "local"
        return jsonify(local)
    try:
        payload = dict(stc_live.chat_context(ROOT, room_id, at, minutes, msg_id=msg_id) or {})
        payload["source"] = "scrm"
        return jsonify(payload)
    except Exception:
        logging.exception("拉取群聊上下文失败")
        return jsonify({"error": "暂时取不到聊天记录，请稍后重试"}), 502


@app.route("/api/timeout/events")
def timeout_events_api():
    db = _db_path()
    if not os.path.isfile(db):
        return jsonify({"total": 0, "events": []})
    try:
        limit = max(1, min(500, int(request.args.get("limit") or 200)))
        offset = max(0, int(request.args.get("offset") or 0))
    except ValueError:
        abort(400)
    conn = open_ro(db)
    try:
        payload = stc_live.query_events(
            conn,
            theater=(request.args.get("theater") or "").strip(),
            q=(request.args.get("q") or "").strip(),
            store=(request.args.get("store") or "").strip(),
            rule_id=(request.args.get("rule_id") or "").strip(),
            msg_id=(request.args.get("msg_id") or "").strip()[:120],
            room_id=(request.args.get("room_id") or "").strip(),
            start=g.date_range[0].isoformat(),
            end=g.date_range[1].isoformat(),
            limit=limit, offset=offset,
            status=(request.args.get("status") or "valid").strip(),
        )
    except ValueError:
        abort(400, description="日期格式应为 YYYY-MM-DD")
    finally:
        conn.close()
    return jsonify(payload)


@app.route("/api/group/<room_id>/members")
def group_members_api(room_id):
    if not _ROOM_RE.match(room_id):
        abort(400)
    try:
        return jsonify({"members": stc_live.group_members(ROOT, room_id)})
    except Exception:
        logging.exception("拉取群成员失败")
        return jsonify({"error": "SCRM 群成员接口暂不可用，请稍后重试"}), 502


@app.route("/api/timeout/context")
def timeout_context_api():
    room_id = (request.args.get("room_id") or "").strip()
    at = (request.args.get("at") or "").strip()
    if not _ROOM_RE.match(room_id) or not re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$", at):
        abort(400)
    try:
        minutes = int(request.args.get("minutes") or 30)
    except ValueError:
        minutes = 30
    try:
        msg_id = (request.args.get("msg_id") or "").strip()[:120]
        return jsonify(stc_live.chat_context(ROOT, room_id, at, min(max(minutes, 1), 240), msg_id=msg_id))
    except Exception:
        logging.exception("拉取会话上下文失败")
        return jsonify({"error": "SCRM 会话存档接口暂不可用，请稍后重试"}), 502


@app.route("/api/timeout/import", methods=["POST"])
@app.route("/api/timeout/import-scan", methods=["POST"])
def timeout_import_gone():
    return jsonify({"error": "Excel 导入已停用：超时改由会话存档自动计算"}), 410


@app.route("/admin/rules")
def admin_rules_page():
    if not _is_admin():
        abort(403)
    return render_template("admin_rules.html", nav_active="admin")


@app.route("/api/admin/rules")
def admin_rules_list():
    if not _is_admin():
        abort(403)
    conn = timeout_algo.open_db(ROOT)
    try:
        rules = timeout_algo.list_rules(conn)
        from timeout_store import get_meta
        last = get_meta(conn, "last_run")
        days = conn.execute("SELECT MIN(day), MAX(day), COUNT(*) FROM chat_msg_day").fetchone()
    finally:
        conn.close()
    return jsonify({"rules": rules, "job": timeout_algo.job_status(),
                    "last_run": json.loads(last) if last else None,
                    "archive": {"from": days[0], "to": days[1], "days": days[2]}})


def _start_recompute(start):
    started = timeout_algo.start_background(ROOT, start=start, fetch=False, label="规则变更重算")
    return started


@app.route("/api/admin/rules", methods=["POST"])
@app.route("/api/admin/rules/<rule_key>", methods=["POST"])
def admin_rules_save(rule_key=None):
    if not _is_admin():
        abort(403)
    data = request.get_json(silent=True) or {}
    conn = timeout_algo.open_db(ROOT)
    try:
        key, since = timeout_algo.save_rule(conn, data, rule_key=rule_key, actor=_actor())
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except KeyError:
        return jsonify({"error": "规则不存在"}), 404
    finally:
        conn.close()
    started = _start_recompute(since)
    return jsonify({"rule_key": key, "recompute_from": since.isoformat(), "started": started,
                    "job": timeout_algo.job_status()})


@app.route("/api/admin/recompute", methods=["POST"])
def admin_recompute():
    if not _is_admin():
        abort(403)
    data = request.get_json(silent=True) or {}
    try:
        start = date.fromisoformat(str(data.get("start"))[:10]) if data.get("start") else None
    except ValueError:
        return jsonify({"error": "日期格式应为 YYYY-MM-DD"}), 400
    started = timeout_algo.start_background(ROOT, start=start, fetch=bool(data.get("fetch")), label="手动重算")
    return jsonify({"started": started, "job": timeout_algo.job_status()})


@app.route("/api/admin/job")
def admin_job():
    if not _is_admin():
        abort(403)
    return jsonify(timeout_algo.job_status())


@app.route("/api/timeout/export")
def timeout_export_csv():
    payload = query_chart(
        ROOT,
        scope=request.args.get("scope") or "national",
        ident=request.args.get("id") or "",
        grain=request.args.get("grain") or "month",
        start=g.date_range[0].isoformat(),
        end=g.date_range[1].isoformat(),
        theater=request.args.get("theater") or "",
        q=request.args.get("q") or "",
    )
    rules = payload.get("totals", {}).get("rules") or []
    output = io.StringIO()
    writer = csv.writer(output)
    header = ["排名", "门店", "战区", "计入消息数", "触发群数", "规则触发次数"]
    for rule in rules:
        header.append((rule.get("rule_name") or rule.get("rule_id")) + "次数")
    writer.writerow(header)
    for idx, row in enumerate(payload.get("stores") or [], 1):
        line = [
            idx,
            row.get("store") or "",
            row.get("theater") or "",
            row.get("messages") or 0,
            row.get("rooms") or 0,
            row.get("triggers") or 0,
        ]
        by_rule = row.get("by_rule") or {}
        for rule in rules:
            cell = by_rule.get(rule.get("rule_id")) or {}
            line.append(cell.get("triggers") or 0)
        writer.writerow(line)
    filename = "超时提醒门店统计.csv"
    return Response(
        output.getvalue().encode("utf-8-sig"),
        mimetype="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename)},
    )


# --- 群名关键字统计（基于 chat_store，只读查询） ---
def _keyword_payload():
    keyword = request.args.get("keyword") or ""
    theater = (request.args.get("theater") or "").strip()
    store_q = (request.args.get("q") or "").strip()
    only_robot = (request.args.get("only_robot") or "0").strip().lower() in ("1", "true", "yes", "on")
    db_path = _db_path()
    if not os.path.isfile(db_path):
        return None, (jsonify({"error": "暂无群数据（找不到 timeout.db），请先从 SCRM 同步"}), 400)
    conn = open_ro(db_path)
    try:
        payload = query_keyword_groups(
            conn, keyword, theater=theater, only_robot=only_robot, store_q=store_q
        )
    except KeywordError as exc:
        return None, (jsonify({"error": str(exc)}), 400)
    finally:
        conn.close()
    if not payload.get("synced_at"):
        data = load_data()
        if data and data.get("updated_at"):
            payload["synced_at"] = data["updated_at"]
    return payload, None


@app.route("/groups")
def groups_page():
    return render_template("groups.html")


@app.route("/api/groups/list")
def api_groups_list():
    page_raw = request.args.get("page", "1")
    size_raw = request.args.get("size", "50")
    try:
        page = int(page_raw if page_raw not in ("", None) else 1)
        size = int(size_raw if size_raw not in ("", None) else 50)
    except (TypeError, ValueError):
        abort(400, description="page/size 必须是正整数")
    if page < 1 or size < 1:
        abort(400, description="page/size 必须是正整数")
    size = min(size, 100)
    robot = (request.args.get("robot") or "all").strip().lower()
    if robot not in ("all", "1", "0"):
        abort(400, description="robot 只能是 all、1 或 0")

    db_path = _db_path()
    if not os.path.isfile(db_path):
        return jsonify({"total": 0, "page": page, "size": size, "rows": [], "theaters": []})
    conn = open_ro(db_path)
    try:
        payload = query_group_list(
            conn,
            q=request.args.get("q") or "",
            theater=request.args.get("theater") or "",
            store=request.args.get("store") or "",
            robot=robot,
            page=page,
            size=size,
        )
    except GroupListError as exc:
        abort(400, description=str(exc))
    finally:
        conn.close()
    return jsonify(payload)


@app.route("/api/groups/keyword")
def api_groups_keyword():
    payload, err = _keyword_payload()
    if err:
        return err
    return jsonify(payload)


@app.route("/api/groups/keyword/export")
def api_groups_keyword_export():
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    payload, err = _keyword_payload()
    if err:
        return err

    wb = Workbook()

    def sheet_rows(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for row in rows:
            ws.append(row)
        if widths:
            for idx, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(idx)].width = width
        return ws

    ws = wb.active
    ws.title = "汇总"
    summary_rows = [
        ("关键字", payload["keyword"]),
        ("战区筛选", payload.get("theater") or "全部"),
        ("门店筛选", payload.get("q") or "全部"),
        ("仅已拉机器人群", "是" if payload.get("only_robot") else "否"),
        ("匹配群总数", payload["total"]),
        ("已拉机器人群数", payload["robot_total"]),
        ("涉及门店数", payload["matched_store_total"]),
        ("未归店群数", payload["unmatched_total"]),
        ("数据同步时间", payload.get("synced_at") or ""),
    ]
    for row in summary_rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 44

    sheet_rows(
        "按战区",
        ["战区", "群数", "门店数"],
        [(r["theater"], r["group_count"], r["store_count"]) for r in payload["by_theater"]],
        widths=[20, 10, 10],
    )
    sheet_rows(
        "按门店",
        ["排名", "门店", "战区", "群数", "已拉机器人"],
        [
            (idx, r["store"], r["theater"], r["group_count"], r["robot_count"])
            for idx, r in enumerate(payload["by_store"], 1)
        ],
        widths=[8, 34, 16, 10, 14],
    )
    sheet_rows(
        "未归店群",
        ["序号", "群名", "是否已拉机器人"],
        [(idx, r["group_name"], "是" if r["in_robot"] else "否") for idx, r in enumerate(payload["unmatched"], 1)],
        widths=[8, 60, 16],
    )
    sheet_rows(
        "群明细",
        ["序号", "群名", "门店", "战区", "是否已拉机器人"],
        [
            (idx, r["group_name"], r["store"], r["theater"], "是" if r["in_robot"] else "否")
            for idx, r in enumerate(payload["groups"], 1)
        ],
        widths=[8, 60, 34, 16, 16],
    )

    buf = io.BytesIO()
    wb.save(buf)
    safe_kw = re.sub(r'[\\/:*?"<>|\r\n]+', "_", payload["keyword"])[:60] or "空"
    filename = f"群名统计_{safe_kw}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename)},
    )
register_score(app, ROOT)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="门店超时提醒机器人统计控制台")
    parser.add_argument("--port", type=int, default=5001, help="监听端口（默认 5001）")
    args = parser.parse_args()
    app.run(host="0.0.0.0", port=args.port, debug=True)
