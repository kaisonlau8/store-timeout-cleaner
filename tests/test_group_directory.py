"""B3 群目录与数据补齐：迁移、同步写入、群详情、群列表与筛选导出。"""

from __future__ import annotations

import importlib
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
from datetime import date, datetime
from pathlib import Path
from urllib.parse import unquote
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

import timeout_algo
from group_keyword import open_ro, query_group_list
from group_query import local_context
from scrm_sync import sync_from_scrm
from tests.test_coverage_store import _minimal_stats
from timeout_store import connect, replace_context

ROOT = Path(__file__).resolve().parent.parent
STORE = "猛士汽车武汉门店"
SHANGHAI = ZoneInfo("Asia/Shanghai")

GROUPS = [
    {
        "group_id": "r1",
        "group_name": "X700-VIP 车主群A",
        "store": STORE,
        "theater": "武汉特区",
        "create_time": "2026-05-01T00:00:00+08:00",
        "owner": "张三",
        "owner_id": "u1",
        "dept": "猛士汽车/门店/武汉特区/" + STORE,
        "member_count": 10,
    },
    {
        "group_id": "r2",
        "group_name": "x700-vip 车主群B",
        "store": STORE,
        "theater": "武汉特区",
        "create_time": "2026-06-01T00:00:00+08:00",
        "owner": "李四",
        "owner_id": "u2",
        "dept": "猛士汽车/门店/武汉特区/" + STORE,
        "member_count": 20,
    },
    {
        "group_id": "r3",
        "group_name": "X700 交付群",
        "store": "猛士汽车上海门店",
        "theater": "华东战区",
        "create_time": "2026-07-01T00:00:00+08:00",
        "owner": "王五",
        "owner_id": "u3",
        "dept": "猛士汽车/门店/华东战区/猛士汽车上海门店",
        "member_count": 30,
    },
    {
        "group_id": "r4",
        "group_name": "VIP 保养群",
        "store": "猛士汽车北京门店",
        "theater": "华北战区",
        "create_time": "2026-04-01T00:00:00+08:00",
        "owner": "赵六",
        "owner_id": "u4",
        "dept": "猛士汽车/门店/华北战区/猛士汽车北京门店",
        "member_count": 40,
    },
    {
        "group_id": "r5",
        "group_name": "未归店群",
        "store": "",
        "theater": "",
        "create_time": None,
    },
    {
        "group_id": "r6",
        "group_name": "X700 上海二群",
        "store": "猛士汽车上海门店",
        "theater": "华东战区",
        "create_time": "2026-08-01T00:00:00+08:00",
        "owner": "钱七",
        "owner_id": "u7",
        "dept": "猛士汽车/门店/华东战区/猛士汽车上海门店",
        "member_count": 0,
    },
]
ROBOT_IDS = ["r1", "r3", "r6"]
EXPORT_HEADERS = ("群名称", "群 ID", "店名", "战区", "是否拉了机器人", "群主", "群人数", "创建时间")


def _write_data(root: Path) -> Path:
    path = root / "data.json"
    path.write_text(json.dumps(_minimal_stats(), ensure_ascii=False), encoding="utf-8")
    return path


def _load_app(monkeypatch, db: Path, data: Path):
    monkeypatch.setenv("CLEANER_DB_PATH", str(db))
    monkeypatch.setenv("CLEANER_DATA_FILE", str(data))
    monkeypatch.delenv("MHERO_AUTH", raising=False)
    monkeypatch.delenv("MHERO_AUTH_CLEANER", raising=False)
    import app as app_module

    return importlib.reload(app_module)


@pytest.fixture
def directory_client(tmp_path, monkeypatch):
    db = tmp_path / "timeout.db"
    conn = connect(str(db))
    try:
        replace_context(
            conn,
            robot_chat_ids=ROBOT_IDS,
            employee_ids=["u1", "u2", "u3", "u4", "u7"],
            group_rows=GROUPS,
            employee_names={"u1": "张三", "u2": "李四"},
        )
        conn.commit()
    finally:
        conn.close()
    module = _load_app(monkeypatch, db, _write_data(tmp_path))
    return module.app.test_client(), module, tmp_path


def _columns(conn, table):
    return {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}


def test_migration_is_idempotent_and_old_rows_have_null_new_columns(tmp_path):
    db = tmp_path / "legacy.db"
    raw = sqlite3.connect(db)
    try:
        raw.executescript(
            """
            CREATE TABLE chat_store (
              room_id TEXT PRIMARY KEY,
              store TEXT,
              theater TEXT,
              group_name TEXT,
              in_robot INTEGER NOT NULL DEFAULT 0
            );
            CREATE TABLE employee (userid TEXT PRIMARY KEY);
            INSERT INTO chat_store(room_id, store, theater, group_name, in_robot)
            VALUES ('old-room', '猛士汽车武汉门店', '武汉特区', '旧群', 1);
            INSERT INTO employee(userid) VALUES ('old-emp');
            """
        )
        raw.commit()
    finally:
        raw.close()

    for _ in range(2):
        conn = connect(str(db))
        try:
            assert {"owner_id", "owner_name", "owner_dept", "member_count", "create_time"} <= _columns(conn, "chat_store")
            assert "name" in _columns(conn, "employee")
            row = conn.execute("SELECT * FROM chat_store WHERE room_id = 'old-room'").fetchone()
            assert row["store"] == "猛士汽车武汉门店"
            assert row["owner_id"] is None
            assert row["owner_name"] is None
            assert row["owner_dept"] is None
            assert row["member_count"] is None
            assert conn.execute("SELECT name FROM employee WHERE userid = 'old-emp'").fetchone()["name"] is None
        finally:
            conn.close()

    conn = connect(str(db))
    try:
        group_cols = [row[1] for row in conn.execute("PRAGMA table_info(chat_store)")]
        employee_cols = [row[1] for row in conn.execute("PRAGMA table_info(employee)")]
        assert group_cols.count("owner_id") == 1
        assert group_cols.count("member_count") == 1
        assert employee_cols.count("name") == 1
        assert "customer_count" not in group_cols
    finally:
        conn.close()


def test_replace_context_writes_group_columns_and_employee_names(tmp_path):
    conn = connect(str(tmp_path / "timeout.db"))
    try:
        replace_context(
            conn,
            robot_chat_ids=["r1"],
            employee_ids=["u1", "u2"],
            employee_names={"u1": "张三", "u2": None},
            group_rows=[
                {
                    "group_id": "r1",
                    "group_name": "甲群",
                    "store": STORE,
                    "theater": "武汉特区",
                    "create_time": "2026-05-01T00:00:00+08:00",
                    "owner": "张三",
                    "owner_id": "u1",
                    "dept": "猛士汽车/门店/武汉特区/" + STORE,
                    "member_count": 36,
                    "customer_count": 0,  # 假数据，不应落入数据库
                },
                {"group_id": "r2", "group_name": "乙群", "store": "", "theater": ""},
            ],
        )
        conn.commit()
        row = conn.execute("SELECT * FROM chat_store WHERE room_id = 'r1'").fetchone()
        assert row["owner_id"] == "u1"
        assert row["owner_name"] == "张三"
        assert row["owner_dept"] == "猛士汽车/门店/武汉特区/" + STORE
        assert row["member_count"] == 36
        assert "customer_count" not in _columns(conn, "chat_store")

        missing = conn.execute("SELECT * FROM chat_store WHERE room_id = 'r2'").fetchone()
        assert missing["owner_id"] is None
        assert missing["owner_name"] is None
        assert missing["owner_dept"] is None
        assert missing["member_count"] is None

        employees = {row["userid"]: row["name"] for row in conn.execute("SELECT userid, name FROM employee")}
        assert employees == {"u1": "张三", "u2": None}
    finally:
        conn.close()


def test_data_json_does_not_persist_employee_names_or_rows(tmp_path, monkeypatch):
    db = tmp_path / "timeout.db"
    data = tmp_path / "data.json"
    module = _load_app(monkeypatch, db, data)
    rows = [dict(GROUPS[0])]
    result = {
        "all_rows": rows,
        "filter_rows": rows,
        "summary": {
            "source": "scrm",
            "all_total": 1,
            "all_valid": 1,
            "all_invalid": 0,
            "filter_total": 1,
            "filter_valid": 1,
            "filter_invalid": 0,
            "robots": [],
            "all_rows": rows,
            "employee_names": {"u1": "张三"},
        },
        "robot_chat_ids": ["r1"],
        "employee_ids": ["u1"],
        "employee_names": {"u1": "张三"},
    }
    monkeypatch.setattr(module, "sync_from_scrm", lambda root: result)

    module._run_group_and_timeout()

    text = data.read_text(encoding="utf-8")
    assert "employee_names" not in text
    assert "all_rows" not in text
    assert "filter_rows" not in text
    payload = json.loads(text)
    assert set(payload) == {"updated_at", "stats", "summary"}

    conn = connect(str(db))
    try:
        assert conn.execute("SELECT name FROM employee WHERE userid = 'u1'").fetchone()["name"] == "张三"
        assert conn.execute("SELECT owner_name FROM chat_store WHERE room_id = 'r1'").fetchone()["owner_name"] == "张三"
    finally:
        conn.close()


def test_sync_from_scrm_returns_employee_names(monkeypatch):
    class FakeClient:
        def list_users(self):
            return [
                {"userId": "u1", "username": "张三", "main_department_id": 4},
                {"userId": "u2", "username": "李四", "main_department_id": 4},
            ]

        def list_departments(self):
            return [
                {"id": 1, "name": "猛士汽车", "parent_id": 0},
                {"id": 2, "name": "门店", "parent_id": 1},
                {"id": 3, "name": "武汉特区", "parent_id": 2},
                {"id": 4, "name": STORE, "parent_id": 3},
            ]

        def list_groups(self, userid=None, user_type=None):
            if userid:
                return [{"chat_id": "g1"}]
            return [
                {
                    "chat_id": "g1",
                    "name": "甲群",
                    "owner": "u1",
                    "create_time": 1777593600,
                    "count": 12,
                }
            ]

    monkeypatch.setattr("scrm_sync.make_scrm_client", lambda root: FakeClient())
    monkeypatch.setenv("SCRM_ROBOT_USERIDS", "u2")
    result = sync_from_scrm(str(ROOT))

    assert result["employee_names"] == {"u1": "张三", "u2": "李四"}
    assert result["all_rows"][0]["owner"] == "张三"
    assert result["all_rows"][0]["owner_id"] == "u1"
    assert result["all_rows"][0]["dept"] == f"猛士汽车/门店/武汉特区/{STORE}"
    assert result["all_rows"][0]["member_count"] == 12


def _seed_detail_page(root: Path) -> None:
    timeout_algo.open_db(str(root)).close()
    conn = connect(str(root / "timeout.db"))
    try:
        replace_context(
            conn,
            robot_chat_ids=["room-owner"],
            employee_ids=["u1"],
            employee_names={"u1": "张三"},
            group_rows=[
                {
                    "group_id": "room-owner",
                    "group_name": "有群主群",
                    "store": STORE,
                    "theater": "武汉特区",
                    "create_time": "2026-05-01T00:00:00+08:00",
                    "owner": "张三",
                    "owner_id": "u1",
                    "dept": "猛士汽车/门店/武汉特区/" + STORE,
                    "member_count": 42,
                },
                {
                    "group_id": "room-empty",
                    "group_name": "未同步群",
                    "store": STORE,
                    "theater": "武汉特区",
                    "create_time": "2026-05-01T00:00:00+08:00",
                },
            ],
        )
        conn.commit()
    finally:
        conn.close()


@pytest.fixture
def detail_client(tmp_path, monkeypatch):
    _seed_detail_page(tmp_path)
    module = _load_app(monkeypatch, tmp_path / "timeout.db", _write_data(tmp_path))
    return module.app.test_client(), module


def test_group_page_renders_owner_department_members_and_unsynced(detail_client):
    client, _ = detail_client
    html = client.get("/group/room-owner?start=2026-08-01&end=2026-08-01").get_data(as_text=True)
    assert "群主 张三（" + STORE + "）" in html
    assert "群人数 42" in html
    assert "客户数" not in html

    empty_html = client.get("/group/room-empty?start=2026-08-01&end=2026-08-01").get_data(as_text=True)
    assert "群主 未同步" in empty_html
    assert "群人数 未同步" in empty_html


def test_local_context_sender_uses_employee_name_and_role_labels(tmp_path):
    root = str(tmp_path)
    conn = timeout_algo.open_db(root)
    try:
        replace_context(
            conn,
            robot_chat_ids=["robot1"],
            employee_ids=["emp1", "emp2"],
            employee_names={"emp1": "张三"},
            group_rows=[],
        )
        base = int(datetime(2026, 8, 1, 8, 50, 0).replace(tzinfo=SHANGHAI).timestamp())
        conn.executemany(
            "INSERT INTO chat_msg(msg_id, room_id, from_id, msgtime, msgtype, content) VALUES (?, ?, ?, ?, ?, ?)",
            [
                ("m1", "room1", "emp1", base, "text", "员工有名字"),
                ("m2", "room1", "emp2", base + 1, "text", "员工没名字"),
                ("m3", "room1", "wm-customer", base + 2, "text", "客户消息"),
                ("m4", "room1", "robot1", base + 3, "text", "机器人消息"),
            ],
        )
        conn.commit()
        payload = local_context(
            conn,
            "room1",
            "m3",
            "2026-08-01 09:00:00",
            30,
            ["emp1", "emp2"],
            ["robot1"],
        )
    finally:
        conn.close()

    by_text = {msg["text"]: msg for msg in payload["messages"]}
    assert by_text["员工有名字"]["sender"] == "张三"
    assert by_text["员工没名字"]["sender"] == "员工"
    assert by_text["客户消息"]["sender"] == "客户"
    assert by_text["机器人消息"]["sender"] == "机器人"
    assert all("sender" in msg for msg in payload["messages"])


def test_groups_list_defaults_sorting_and_shape(directory_client):
    client, _, _ = directory_client
    response = client.get("/api/groups/list")
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["total"] == 6
    assert payload["page"] == 1
    assert payload["size"] == 50
    assert payload["theaters"] == ["华东战区", "华北战区", "武汉特区"]
    assert [row["room_id"] for row in payload["rows"]] == ["r6", "r3", "r2", "r1", "r4", "r5"]
    assert set(payload["rows"][0]) == {
        "room_id", "group_name", "store", "theater", "in_robot",
        "owner_name", "member_count", "create_time",
    }
    missing = next(row for row in payload["rows"] if row["room_id"] == "r5")
    assert missing["owner_name"] is None
    assert missing["member_count"] is None
    assert missing["create_time"] == ""
    assert missing["store"] == ""


@pytest.mark.parametrize(
    "query,total,ids",
    [
        ("q=x700", 4, ["r6", "r3", "r2", "r1"]),
        ("q=R2", 1, ["r2"]),
        ("theater=华东战区", 2, ["r6", "r3"]),
        ("store=上海", 2, ["r6", "r3"]),
        ("robot=1", 3, ["r6", "r3", "r1"]),
        ("robot=0", 3, ["r2", "r4", "r5"]),
        ("theater=华东战区&robot=0", 0, []),
        ("theater=华东战区&robot=1&store=上海&q=x700", 2, ["r6", "r3"]),
    ],
)
def test_groups_list_filters(directory_client, query, total, ids):
    client, _, _ = directory_client
    payload = client.get("/api/groups/list?" + query).get_json()
    assert payload["total"] == total
    assert [row["room_id"] for row in payload["rows"]] == ids


def test_groups_list_pagination_size_cap_and_invalid_params(directory_client):
    client, _, _ = directory_client
    page2 = client.get("/api/groups/list?page=2&size=2").get_json()
    assert page2["page"] == 2
    assert page2["size"] == 2
    assert page2["total"] == 6
    assert [row["room_id"] for row in page2["rows"]] == ["r2", "r1"]

    capped = client.get("/api/groups/list?size=999").get_json()
    assert capped["size"] == 100
    assert capped["total"] == 6

    for query in ("page=0", "page=-1", "page=abc", "size=0", "size=abc", "robot=2"):
        response = client.get("/api/groups/list?" + query)
        assert response.status_code == 400, query


def _read_export(response):
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    disposition = unquote(response.headers["Content-Disposition"])
    assert re.fullmatch(r"attachment; filename\*=UTF-8''群列表_\d{8}\.xlsx", disposition)
    wb = load_workbook(io.BytesIO(response.data), read_only=True)
    try:
        assert wb.sheetnames == ["群列表"]
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[0] == EXPORT_HEADERS
        return [dict(zip(EXPORT_HEADERS, row)) for row in rows[1:]]
    finally:
        wb.close()


@pytest.mark.parametrize(
    "filters,ids",
    [
        ({}, ["r6", "r3", "r2", "r1", "r4", "r5"]),
        ({"theater": "华东战区"}, ["r6", "r3"]),
        ({"theater": "武汉特区"}, ["r2", "r1"]),
        ({"robot": "0"}, ["r2", "r4", "r5"]),
        ({"robot": "1"}, ["r6", "r3", "r1"]),
        ({"store": "上海"}, ["r6", "r3"]),
        ({"q": "x700"}, ["r6", "r3", "r2", "r1"]),
        ({"q": "R2"}, ["r2"]),
        ({"q": "不存在的群"}, []),
        ({"theater": "华东战区", "robot": "0"}, []),
        ({"theater": " 华东战区 ", "robot": " 1 ", "store": " 上海 ", "q": " x700 "}, ["r6", "r3"]),
    ],
)
def test_groups_list_export_matches_query_filters_total_and_content(directory_client, filters, ids):
    client, _, root = directory_client
    conn = open_ro(str(root / "timeout.db"))
    try:
        listed = query_group_list(conn, **filters)
    finally:
        conn.close()

    exported = _read_export(client.get(
        "/api/groups/list/export", query_string={**filters, "page": "3", "size": "1"}
    ))
    assert len(exported) == listed["total"] == len(ids)
    assert [row["群 ID"] for row in exported] == ids
    for actual, source in zip(exported, listed["rows"]):
        assert actual == {
            "群名称": source["group_name"] or None,
            "群 ID": source["room_id"],
            "店名": source["store"] or None,
            "战区": source["theater"] or None,
            "是否拉了机器人": "是" if source["in_robot"] else "否",
            "群主": source["owner_name"],
            "群人数": source["member_count"],
            "创建时间": source["create_time"] or None,
        }


def test_groups_list_export_exceeds_page_cap_ignores_pagination_and_sorts_ties(directory_client):
    client, _, root = directory_client
    ids = [f"bulk-{idx:03d}" for idx in range(125)]
    conn = connect(str(root / "timeout.db"))
    try:
        conn.executemany(
            "INSERT INTO chat_store(room_id, group_name, in_robot, create_time) VALUES (?, ?, 0, ?)",
            [(room_id, "批量测试群", "2026-09-01T00:00:00+08:00") for room_id in reversed(ids)],
        )
        conn.commit()
    finally:
        conn.close()

    listed = client.get("/api/groups/list?q=bulk-&page=2&size=50").get_json()
    assert listed["total"] == 125
    assert len(listed["rows"]) == 50
    assert [row["room_id"] for row in listed["rows"]] == ids[50:100]

    exported = _read_export(client.get("/api/groups/list/export?q=bulk-&page=abc&size=0"))
    assert len(exported) == listed["total"]
    assert [row["群 ID"] for row in exported] == ids


@pytest.mark.parametrize("missing_db", [True, False])
def test_groups_list_export_empty_or_missing_database_returns_headers(tmp_path, monkeypatch, missing_db):
    db = tmp_path / "timeout.db"
    if not missing_db:
        connect(str(db)).close()
    module = _load_app(monkeypatch, db, _write_data(tmp_path))
    client = module.app.test_client()
    assert _read_export(client.get("/api/groups/list/export?robot=0")) == []
    assert client.get("/api/groups/list?robot=0").get_json()["total"] == 0
    assert db.exists() is not missing_db


@pytest.mark.parametrize("missing_db", [True, False])
@pytest.mark.parametrize("robot", ["2", "true", "-1", " "])
def test_groups_list_export_invalid_robot_matches_list_400(directory_client, monkeypatch, missing_db, robot):
    client, _, root = directory_client
    if missing_db:
        monkeypatch.setenv("CLEANER_DB_PATH", str(root / "missing.db"))
    listed = client.get("/api/groups/list", query_string={"robot": robot})
    exported = client.get("/api/groups/list/export", query_string={"robot": robot})
    assert exported.status_code == listed.status_code == 400
    assert exported.data == listed.data
    assert "robot 只能是 all、1 或 0" in exported.get_data(as_text=True)


def test_groups_list_export_preserves_formula_like_text_and_numeric_zero(directory_client):
    client, _, root = directory_client
    conn = connect(str(root / "timeout.db"))
    try:
        conn.execute("UPDATE chat_store SET group_name = '=1+1', owner_name = '=2+2' WHERE room_id = 'r6'")
        conn.commit()
    finally:
        conn.close()
    response = client.get("/api/groups/list/export?q=r6")
    exported = _read_export(response)
    assert exported[0]["群名称"] == "=1+1"
    assert exported[0]["群主"] == "=2+2"
    assert exported[0]["群人数"] == 0
    wb = load_workbook(io.BytesIO(response.data))
    try:
        assert wb.active["A2"].data_type == "s"
        assert wb.active["F2"].data_type == "s"
    finally:
        wb.close()


def test_groups_page_has_tabs_list_table_and_management_title(directory_client):
    client, _, _ = directory_client
    html = client.get("/groups").get_data(as_text=True)
    assert "<title>群管理 · 超时机器人统计</title>" in html
    assert "<h1>群管理</h1>" in html
    assert "群名统计" not in html
    assert 'data-tab="list"' in html and 'data-tab="keyword"' in html
    assert 'class="tab-content active" id="tab-list"' in html
    assert 'class="tab-content" id="tab-keyword"' in html
    assert 'id="keyword-input"' in html and 'id="store-q"' in html
    assert 'id="groups-list-table"' in html and "data-ds-primary" in html
    assert 'class="export-link groups-export" id="groups-list-export" href="/api/groups/list/export"' in html
    assert html.index('id="groups-list-query"') < html.index('id="groups-list-export"') < html.index('id="groups-list-table"')
    for header in ("群名", "门店", "战区", "机器人", "群主", "群人数", "创建时间"):
        assert "<th" in html and header in html
    assert "js/groups-list.js" in html
    assert "js/antd-feedback.js" in html

    keyword = client.get("/groups?tab=keyword").get_data(as_text=True)
    assert 'class="tab-content" id="tab-list"' in keyword
    assert 'class="tab-content active" id="tab-keyword"' in keyword

    nav = (ROOT / "templates" / "_nav.html").read_text(encoding="utf-8")
    assert "群管理" in nav and "群名统计" not in nav


def test_groups_list_js_uses_spin_replace_state_and_links(directory_client):
    source = (ROOT / "static" / "js" / "groups-list.js").read_text(encoding="utf-8")
    assert "AntdFeedback.setSpin" in source
    assert "history.replaceState" in source
    assert "params.set('tab'" in source
    assert '"/group/' in source and "encodeURIComponent(roomId)" in source
    assert '"/store/' in source and "encodeURIComponent(row.store)" in source
    assert "未同步" in source and "未归店" in source
    assert "上一页" in (ROOT / "templates" / "groups.html").read_text(encoding="utf-8")


def test_groups_list_export_js_uses_current_filters_without_pagination():
    node = shutil.which("node")
    if not node:
        pytest.skip("node not installed")
    result = subprocess.run(
        [node, str(ROOT / "tests" / "groups_list_export_js.cjs")],
        cwd=ROOT, capture_output=True, text=True, timeout=10,
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "ok"


def test_groups_list_responsive_fixes_are_scoped_and_versioned(directory_client):
    client, _, _ = directory_client
    html = client.get("/groups").get_data(as_text=True)
    assert "css/groups-list.css?v=2" in html
    assert "js/groups-list.js?v=6" in html
    assert "js/groups.js?v=8" in html

    css = (ROOT / "static" / "css" / "groups-list.css").read_text(encoding="utf-8")
    assert "#groups-list-table-wrap" in css
    assert "max-height: none" in css
    assert "overflow: auto" in css
    assert "flex-wrap: nowrap" in css
    assert "text-overflow: ellipsis" in css
    # 卡片「标签左、值右」由设计系统统一处理，应用不再局部覆盖卡片单元格
    assert "data-label" not in css
    assert "data-cards-full" not in css
    assert "display: grid !important" not in css

    source = (ROOT / "static" / "js" / "groups-list.js").read_text(encoding="utf-8")
    assert "(max-width: 860px)" in source
    assert "页 · ' + total + ' 个群" in source

    groups_js = (ROOT / "static" / "js" / "groups.js").read_text(encoding="utf-8")
    assert "data-cards-full" not in groups_js
