import os
import tempfile
import unittest
from datetime import date

from openpyxl import Workbook, load_workbook

from score_export import stores_to_csv, stores_to_xlsx
from score_query import build_report, last_quarter, store_detail
from timeout_excel import _msg_id
from timeout_store import connect


def _seed(root: str) -> None:
    os.makedirs(os.path.join(root, "input_file"), exist_ok=True)
    conn = connect(os.path.join(root, "timeout.db"))
    try:
        groups = [
            ("r1", "甲店", "华东战区", "甲店客户群1", 1),
            ("r2", "甲店", "华东战区", "甲店客户群2", 1),
            ("r3", "未接入店", "华北战区", "未接入客户群", 0),
            ("r4", "猛士汽车杭州西溪零售中心", "华东战区", "西溪零售群", 1),
            ("r5", "猛士杭州西溪交付服务中心", "华东战区", "西溪交付群", 0),
        ]
        conn.executemany(
            "INSERT INTO chat_store(room_id, store, theater, group_name, in_robot) VALUES (?, ?, ?, ?, ?)",
            groups,
        )
        events = [
            ("2026-08-01 09:00:00", "甲店客户群1", "甲店", "刹车有异响", "规则10分钟"),
            ("2026-08-01 09:05:00", "甲店客户群1", "甲店", "好的", "规则10分钟"),
            ("2026-08-02 11:00:00", "西溪零售群", "猛士汽车杭州西溪零售中心", "想预约保养", "规则10分钟"),
            ("2026-06-01 11:00:00", "甲店客户群1", "甲店", "季度外的消息", "规则10分钟"),
            ("2026-08-03 12:00:00", "未知群", "", "对不上门店", "规则10分钟"),
        ]
        for created_at, room, store, msg, rule in events:
            conn.execute(
                """
                INSERT INTO timeout_event(
                  msg_id, rule_id, rule_name, room_id, store, theater, created_at, customer_id, source, room_name, content
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    _msg_id(created_at, room, msg),
                    rule,
                    rule,
                    "name:" + room,
                    store,
                    "",
                    created_at,
                    "",
                    "api",
                    room,
                    msg,
                ),
            )
        conn.commit()
    finally:
        conn.close()

    book = Workbook()
    sheet = book.active
    sheet.append(["触发时间", "所在会话", "触发超时规则", "超时消息", "会话类型"])
    sheet.append(["2026-08-01 09:00:00", "甲店客户群1", "规则10分钟", "刹车有异响", "外部群聊"])
    sheet.append(["2026-08-01 09:05:00", "甲店客户群1", "规则10分钟", "好的", "外部群聊"])
    sheet.append(["2026-08-02 11:00:00", "西溪零售群", "规则10分钟", "想预约保养", "外部群聊"])
    sheet.append(["2026-06-01 11:00:00", "甲店客户群1", "规则10分钟", "季度外的消息", "外部群聊"])
    sheet.append(["2026-08-03 12:00:00", "未知群", "规则10分钟", "对不上门店", "外部群聊"])
    book.save(os.path.join(root, "input_file", "导出时间范围20260701-20260930回复超时记录.xlsx"))


class ScoreQueryTests(unittest.TestCase):
    def test_last_quarter(self):
        self.assertEqual(last_quarter(date(2026, 10, 16)), (date(2026, 7, 1), date(2026, 9, 30)))
        self.assertEqual(last_quarter(date(2026, 1, 15)), (date(2025, 10, 1), date(2025, 12, 31)))
        self.assertEqual(last_quarter(date(2026, 4, 1)), (date(2026, 1, 1), date(2026, 3, 31)))

    def test_report_filters_range_and_special_cases(self):
        with tempfile.TemporaryDirectory() as root:
            _seed(root)
            report = build_report(root, date(2026, 7, 1), date(2026, 9, 30))
            by = {row["name"]: row for row in report["stores"]}
            self.assertEqual(by["甲店"]["raw"], 2)
            self.assertEqual(by["甲店"]["valid"], 1)
            self.assertEqual(by["甲店"]["robot_group_count"], 2)
            self.assertEqual(by["甲店"]["theater"], "华东战区")
            self.assertEqual(by["未接入店"]["kind"], "zero")
            self.assertEqual(by["未接入店"]["score"], 0)
            delivery = by["猛士杭州西溪交付服务中心"]
            self.assertEqual(delivery["kind"], "merge")
            self.assertEqual(delivery["valid"], 1)
            self.assertEqual(delivery["robot_group_count"], 1)
            self.assertEqual(report["unmatched"], 1)
            self.assertNotIn("季度外", "".join(row["name"] for row in report["stores"]))

    def test_detail_splits_counted_and_excluded_and_shares_center(self):
        with tempfile.TemporaryDirectory() as root:
            _seed(root)
            detail = store_detail(root, "甲店", date(2026, 7, 1), date(2026, 9, 30))
            self.assertEqual(len(detail["counted"]), 1)
            self.assertEqual(detail["counted"][0]["message"], "刹车有异响")
            self.assertEqual(len(detail["excluded"]), 1)
            self.assertIn("寒暄", detail["excluded"][0]["reason"])
            shared = store_detail(root, "猛士杭州西溪交付服务中心", date(2026, 7, 1), date(2026, 9, 30))
            self.assertEqual(len(shared["counted"]), 1)
            self.assertTrue(shared["counted"][0]["shared"])
            self.assertEqual(shared["counted"][0]["message"], "想预约保养")

    def test_export_csv_and_xlsx_roundtrip(self):
        with tempfile.TemporaryDirectory() as root:
            _seed(root)
            report = build_report(root, date(2026, 7, 1), date(2026, 9, 30))
            csv_text = stores_to_csv(report["stores"]).decode("utf-8-sig")
            self.assertIn("门店名称", csv_text)
            self.assertIn("机器人群数", csv_text)
            self.assertIn("甲店", csv_text)
            self.assertNotIn("刹车有异响", csv_text)
            blob = stores_to_xlsx(report["stores"])
            book = load_workbook(filename=__import__("io").BytesIO(blob))
            sheet = book.active
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            self.assertIn("有效超时", headers)
            names = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
            self.assertIn("甲店", names)


if __name__ == "__main__":
    unittest.main()
