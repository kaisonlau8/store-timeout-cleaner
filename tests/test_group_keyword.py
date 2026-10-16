import json
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

ROOT = Path(__file__).resolve().parent.parent
AUTH_DIR = ROOT.parent / "m-hero-auth"
if AUTH_DIR.is_dir() and str(AUTH_DIR) not in sys.path:
    sys.path.insert(0, str(AUTH_DIR))

from openpyxl import load_workbook

from group_keyword import KeywordError, query_keyword_groups
from timeout_store import connect, replace_context, set_meta

ROWS = [
    {"group_id": "r1", "group_name": "X700-VIP 车主群A", "store": "猛士汽车武汉门店", "theater": "武汉特区"},
    {"group_id": "r2", "group_name": "x700-vip 车主群B", "store": "猛士汽车武汉门店", "theater": "武汉特区"},
    {"group_id": "r3", "group_name": "X700 交付群", "store": "猛士汽车上海门店", "theater": "华东战区"},
    {"group_id": "r4", "group_name": "x700 服务群", "store": "猛士汽车上海门店", "theater": "华东战区"},
    {"group_id": "r5", "group_name": "X700 群主未归店", "store": "", "theater": ""},
    {"group_id": "r6", "group_name": "VIP 保养群", "store": "猛士汽车北京门店", "theater": "华北战区"},
    {"group_id": "r7", "group_name": "X700 无战区", "store": "猛士汽车广州门店", "theater": ""},
]
ROBOT_IDS = ["r1", "r3", "r6"]


def make_conn():
    conn = connect(":memory:")
    replace_context(conn, robot_chat_ids=ROBOT_IDS, employee_ids=["u1"], group_rows=ROWS)
    set_meta(conn, "updated_at", "2026-10-15 00:00:00")
    conn.commit()
    return conn


class QueryKeywordGroupsTests(unittest.TestCase):
    def setUp(self):
        self.conn = make_conn()

    def tearDown(self):
        self.conn.close()

    def test_case_insensitive_substring(self):
        data = query_keyword_groups(self.conn, "  X700-VIP ")
        self.assertEqual(data["keyword"], "X700-VIP")
        self.assertEqual(data["total"], 2)
        self.assertEqual(sorted(g["group_name"] for g in data["groups"]),
                         ["X700-VIP 车主群A", "x700-vip 车主群B"])

    def test_totals_and_robot(self):
        data = query_keyword_groups(self.conn, "x700")
        self.assertEqual(data["total"], 6)  # r1-r5、r7
        self.assertEqual(data["robot_total"], 2)  # r1、r3
        self.assertEqual(data["matched_store_total"], 3)  # 武汉/上海/广州
        self.assertEqual(data["unmatched_total"], 2)  # r5、r7

    def test_theater_filter_applies_to_all(self):
        data = query_keyword_groups(self.conn, "x700", theater="华东战区")
        self.assertEqual(data["total"], 2)
        self.assertEqual([t["theater"] for t in data["by_theater"]], ["华东战区"])
        self.assertEqual([s["store"] for s in data["by_store"]], ["猛士汽车上海门店"])
        self.assertEqual(data["unmatched_total"], 0)
        data_none = query_keyword_groups(self.conn, "x700", theater="不存在战区")
        self.assertEqual(data_none["total"], 0)

    def test_only_robot(self):
        data = query_keyword_groups(self.conn, "x700", only_robot=True)
        names = sorted(g["group_name"] for g in data["groups"])
        self.assertEqual(names, ["X700 交付群", "X700-VIP 车主群A"])
        self.assertEqual(data["robot_total"], data["total"])

    def test_unmatched_list(self):
        data = query_keyword_groups(self.conn, "x700")
        unmatched = {row["room_id"]: row for row in data["unmatched"]}
        self.assertEqual(set(unmatched), {"r5", "r7"})
        self.assertEqual(unmatched["r5"]["group_name"], "X700 群主未归店")
        self.assertEqual(unmatched["r5"]["in_robot"], 0)
        self.assertIn("in_robot", unmatched["r7"])
        self.assertNotIn("store", unmatched["r5"])

    def test_sorting_desc(self):
        data = query_keyword_groups(self.conn, "x700")
        counts = [t["group_count"] for t in data["by_theater"]]
        self.assertEqual(counts, sorted(counts, reverse=True))
        store_counts = [s["group_count"] for s in data["by_store"]]
        self.assertEqual(store_counts, sorted(store_counts, reverse=True))
        self.assertEqual(data["by_store"][0]["store"], "猛士汽车上海门店")
        self.assertEqual(data["by_store"][0]["robot_count"], 1)

    def test_meta_fields(self):
        data = query_keyword_groups(self.conn, "x700")
        self.assertIsNone(data["synced_at"])  # 由 app 从 data.json 填充
        self.assertEqual(data["theaters"], ["华东战区", "华北战区", "武汉特区"])

    def test_store_query_combines_with_theater(self):
        data = query_keyword_groups(self.conn, "x700", theater="华东战区", store_q="上海")
        self.assertEqual(data["q"], "上海")
        self.assertEqual(data["total"], 2)
        self.assertEqual([row["store"] for row in data["by_store"]], ["猛士汽车上海门店"])
        self.assertEqual(data["theaters"], ["华东战区", "华北战区", "武汉特区"])
        none = query_keyword_groups(self.conn, "x700", theater="华东战区", store_q="武汉")
        self.assertEqual(none["total"], 0)
        self.assertEqual(none["by_store"], [])

    def test_empty_keyword_raises(self):
        for bad in ("", "   ", None):
            with self.assertRaises(KeywordError):
                query_keyword_groups(self.conn, bad)


class GroupsApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        db_path = os.path.join(cls.tmp.name, "timeout.db")
        data_file = os.path.join(cls.tmp.name, "data.json")
        conn = connect(db_path)
        replace_context(conn, robot_chat_ids=ROBOT_IDS, employee_ids=["u1"], group_rows=ROWS)
        set_meta(conn, "updated_at", "2026-10-15 00:00:00")
        conn.commit()
        conn.close()
        with open(data_file, "w", encoding="utf-8") as fh:
            json.dump({"updated_at": "2026-10-15 00:00:00", "stats": {}, "summary": {}}, fh)
        os.environ["CLEANER_DB_PATH"] = db_path
        os.environ["CLEANER_DATA_FILE"] = data_file
        os.environ.pop("MHERO_AUTH", None)
        os.environ.pop("MHERO_AUTH_CLEANER", None)
        import app as app_module

        cls.app = app_module.app
        cls.client = app_module.app.test_client()

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        os.environ.pop("CLEANER_DB_PATH", None)
        os.environ.pop("CLEANER_DATA_FILE", None)

    def test_page_ok(self):
        resp = self.client.get("/groups")
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)
        self.assertIn("群名统计", html)
        self.assertIn("门店名称或编码", html)
        self.assertIn('id="store-q"', html)

    def test_api_ok(self):
        resp = self.client.get("/api/groups/keyword?keyword=x700")
        self.assertEqual(resp.status_code, 200)
        payload = resp.get_json()
        self.assertEqual(payload["total"], 6)
        self.assertEqual(payload["robot_total"], 2)
        self.assertEqual(payload["matched_store_total"], 3)
        self.assertEqual(payload["unmatched_total"], 2)
        self.assertEqual(payload["synced_at"], "2026-10-15 00:00:00")

    def test_api_theater_and_robot_params(self):
        resp = self.client.get("/api/groups/keyword?keyword=x700&theater=%E5%8D%8E%E4%B8%9C%E6%88%98%E5%8C%BA&only_robot=1")
        payload = resp.get_json()
        self.assertEqual(payload["total"], 1)  # r3 X700 交付群
        self.assertEqual(payload["groups"][0]["group_name"], "X700 交付群")

    def test_api_store_query_with_theater(self):
        resp = self.client.get("/api/groups/keyword?keyword=x700&theater=%E5%8D%8E%E4%B8%9C%E6%88%98%E5%8C%BA&q=%E4%B8%8A%E6%B5%B7")
        self.assertEqual(resp.status_code, 200)
        payload = resp.get_json()
        self.assertEqual(payload["total"], 2)
        self.assertEqual(payload["q"], "上海")
        self.assertEqual(payload["by_store"][0]["store"], "猛士汽车上海门店")

    def test_api_empty_keyword_400(self):
        resp = self.client.get("/api/groups/keyword?keyword=")
        self.assertEqual(resp.status_code, 400)
        self.assertIn("error", resp.get_json())

    def test_export_xlsx_readable(self):
        resp = self.client.get("/api/groups/keyword/export?keyword=x700")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp.headers["Content-Type"])
        self.assertIn("filename*=UTF-8''", resp.headers["Content-Disposition"])
        wb = load_workbook(io.BytesIO(resp.data))
        self.assertEqual(wb.sheetnames, ["汇总", "按战区", "按门店", "未归店群", "群明细"])
        summary = {row[0]: row[1] for row in wb["汇总"].iter_rows(values_only=True)}
        self.assertEqual(summary["匹配群总数"], 6)
        self.assertEqual(summary["涉及门店数"], 3)
        store_rows = list(wb["按门店"].iter_rows(values_only=True))
        self.assertEqual(store_rows[0], ("排名", "门店", "战区", "群数", "已拉机器人"))
        self.assertEqual(store_rows[1][1], "猛士汽车上海门店")
        detail_rows = list(wb["群明细"].iter_rows(values_only=True))
        self.assertEqual(len(detail_rows), 7)  # 表头 + 6 行


if __name__ == "__main__":
    unittest.main()
