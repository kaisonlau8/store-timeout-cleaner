import copy
import json
import sqlite3
from datetime import datetime, timedelta
from types import SimpleNamespace

import pytest
from flask import Flask, g
from openpyxl import load_workbook

import weekly_report as wr
import weekly_report_state as state
from weekly_report_web import register_weekly_reports

WEEK = "2026-10-19"


def group(at, **extra):
    return {"room_id": at, "create_time": at, "store": "甲店", "theater": "华东", "in_robot": 1, **extra}


def test_week_and_time_boundaries():
    rows = [group(at) for at in [
        "2026-10-18T23:59:59+08:00", "2026-10-18T16:00:00Z",
        "2026-10-25T23:59:59.999999+08:00", "2026-10-26T00:00:00+08:00", "bad", "",
    ]]
    before = copy.deepcopy(rows)
    report = wr.build_weekly_report(WEEK, rows, [])
    assert len(report["new_groups"]) == 2
    assert report["new_robot_groups"] == 2
    assert report["invalid_create_times"] == 2
    assert rows == before
    assert wr.parse_time("2026-10-19 00:00:00").utcoffset() == timedelta(hours=8)
    assert wr.previous_week_start(datetime(2026, 10, 26).date()).isoformat() == WEEK
    with pytest.raises(ValueError):
        wr.week_bounds("2026-10-20")


def test_theaters_and_metrics():
    rows = [group(WEEK), group(WEEK, in_robot="0", store="乙店"), group(WEEK, theater="", store="", in_robot=0)]
    metrics = [
        {"tier": "groups", "start": "2026-10-25T12:00:00+08:00", "status": "failed", "counts": {"groups": 12, "new_groups": 3, "robot_rooms_added": 7, "groups_updated": 8}},
        {"tier": "groups", "start": WEEK, "status": "success", "counts": {"groups": 10, "new_groups": 2, "updated_groups": 4, "groups_updated": 4, "renamed_groups": 2, "in_robot_updated": 1}},
        {"tier": "members", "start": WEEK, "status": "failed", "counts": {"new_groups": 999}},
        {"tier": "groups", "start": "2026-10-26", "status": "failed"},
    ]
    before = copy.deepcopy(metrics)
    report = wr.build_weekly_report(WEEK, rows, metrics)
    assert report["coverage"] == {"groups": 3, "robot_groups": 1, "coverage_rate": 1/3, "stores": 2, "robot_stores": 1}
    assert report["theaters"][0]["new_groups"] == 2
    assert report["sync"]["new_groups"] == 5
    assert report["sync"]["updated_groups"] == 12
    assert report["sync"]["failures"] == 1
    assert report["sync"]["opening_counts"]["groups"] == 10
    assert report["sync"]["closing_counts"]["groups"] == 12
    assert metrics == before


def test_timeout_and_excel(tmp_path):
    events = [{"created_at": WEEK, "theater": "华东", "exclude_reason": value} for value in ("", None, "剔除")]
    events.append({"created_at": "2026-10-26", "exclude_reason": ""})
    report = wr.build_weekly_report(WEEK, [group(WEEK, group_name="=1+1")], [], events)
    assert report["timeout"]["events"] == 3
    assert report["timeout"]["valid"] == 2
    summary = {"events": 3, "valid": 2, "by_theater": [["华东", 3]], "valid_by_theater": [["华东", 2]]}
    assert wr.build_weekly_report(WEEK, [], [], timeout_summary=summary)["timeout"] == report["timeout"]
    summary["events"] = 4
    with pytest.raises(ValueError):
        wr.build_weekly_report(WEEK, [], [], timeout_summary=summary)
    path = wr.write_weekly_excel(report, tmp_path / "test.xlsx")
    book = load_workbook(path)
    try:
        assert book.sheetnames[:4] == ["本周新建群", "战区覆盖", "同步运行", "超时事件"]
        assert book["本周新建群"]["B2"].data_type == "s"
        assert book["本周新建群"]["B2"].value == "=1+1"
    finally:
        book.close()
    brief = wr.render_brief(report, "/weekly-reports/test.xlsx")
    assert "含历史群" in brief and "有效 2" in brief
    assert "未提供数据" in wr.render_brief(wr.build_weekly_report(WEEK, [], []), "url")


@pytest.fixture
def inputs(tmp_path, monkeypatch):
    rows = tmp_path / "rows.json"
    rows.write_text(json.dumps([group(WEEK)]))
    metrics = tmp_path / "metrics.jsonl"
    metrics.write_text("")
    monkeypatch.setattr(wr, "send_chat_text", lambda *args: pytest.fail("dry-run 不得发送"))
    return dict(chat_store_json=rows, metrics_jsonl=metrics, state_path=tmp_path / "state.json", out_dir=tmp_path / "out")


def test_dry_run_then_send_and_dedupe(inputs, monkeypatch):
    result = wr.run_weekly_report(WEEK, dry_run=True, **inputs)
    assert result["dry_run"] and (inputs["out_dir"] / "brief.txt").is_file()
    assert state.load_state(inputs["state_path"])[WEEK]["status"] == "dry_run"
    calls = []
    monkeypatch.setattr(wr, "send_chat_text", lambda *args: calls.append(args) or {"message_id": "mock"})
    wr.run_weekly_report(WEEK, dry_run=False, **inputs)
    assert len(calls) == 1
    assert wr.run_weekly_report(WEEK, dry_run=False, **inputs)["skip_reason"] == "sent"
    wr.run_weekly_report(WEEK, dry_run=True, **inputs)
    assert state.load_state(inputs["state_path"])[WEEK]["status"] == "sent"
    wr.run_weekly_report(WEEK, dry_run=False, force=True, **inputs)
    assert len(calls) == 2


@pytest.mark.parametrize("at,blocked", [("bad", True), ("2099-01-01", True), (None, True), ("2020-01-01", False)])
def test_sending_timeout(inputs, monkeypatch, at, blocked):
    previous = {"status": "sending", "at": at}
    state.save_state(inputs["state_path"], {WEEK: previous})
    monkeypatch.setattr(wr, "send_chat_text", lambda *args: {})
    result = wr.run_weekly_report(WEEK, dry_run=False, **inputs)
    assert result["skipped"] == blocked
    if blocked:
        assert result["skip_reason"] == "sending"
        wr.run_weekly_report(WEEK, dry_run=True, **inputs)
        assert state.load_state(inputs["state_path"])[WEEK] == previous


def test_sending_current_and_send_failure(inputs, monkeypatch):
    previous = {"status": "sending", "at": datetime.now(wr.SHANGHAI).isoformat()}
    state.save_state(inputs["state_path"], {WEEK: previous})
    assert wr.run_weekly_report(WEEK, dry_run=False, **inputs)["skip_reason"] == "sending"

    def fail(*args):
        raise RuntimeError("mock failure")

    monkeypatch.setattr(wr, "send_chat_text", fail)
    with pytest.raises(RuntimeError):
        wr.run_weekly_report(WEEK, dry_run=False, force=True, **inputs)
    assert state.load_state(inputs["state_path"])[WEEK] == previous
    state.save_state(inputs["state_path"], {})
    with pytest.raises(RuntimeError):
        wr.run_weekly_report(WEEK, dry_run=False, **inputs)
    assert state.load_state(inputs["state_path"])[WEEK]["status"] == "failed"


def test_lock_force_and_corruption(inputs):
    with state.report_lock(inputs["state_path"]) as acquired:
        assert acquired
        assert wr.run_weekly_report(WEEK, dry_run=True, force=True, **inputs)["skip_reason"] == "locked"
    inputs["state_path"].write_text('{"week": {"status": "broken"}}')
    with pytest.raises(ValueError):
        wr.run_weekly_report(WEEK, dry_run=True, force=True, **inputs)
    inputs["state_path"].write_text("broken")
    with pytest.raises(ValueError):
        state.load_state(inputs["state_path"])


def test_atomic_write_failure(tmp_path, monkeypatch):
    path = tmp_path / "state.json"
    state.save_state(path, {WEEK: {"status": "sent"}})
    before = path.read_bytes()

    def fail(*args):
        raise OSError("mock replace failure")

    monkeypatch.setattr(state.os, "replace", fail)
    with pytest.raises(OSError):
        state.save_state(path, {})
    assert path.read_bytes() == before
    assert not list(tmp_path.glob("*.tmp"))


def test_sent_write_failure_preserves_sending(inputs, monkeypatch):
    real_save = wr.save_state

    def save(path, data):
        if data[WEEK]["status"] == "sent":
            raise OSError("mock disk failure")
        real_save(path, data)

    monkeypatch.setattr(wr, "save_state", save)
    monkeypatch.setattr(wr, "send_chat_text", lambda *args: {})
    with pytest.raises(OSError):
        wr.run_weekly_report(WEEK, dry_run=False, **inputs)
    assert state.load_state(inputs["state_path"])[WEEK]["status"] == "sending"


@pytest.mark.parametrize("hour,minute", [(0, 0), (23, 55), (8, 35), (9, 5), (9, 25), (9, 35), (12, 4), (24, 0)])
def test_invalid_schedule(tmp_path, hour, minute):
    path = tmp_path / "config.json"
    path.write_text(json.dumps({"hour": hour, "minute": minute}))
    with pytest.raises(ValueError):
        wr.load_config(path)


def test_cli_and_scheduled_disabled():
    with pytest.raises(SystemExit) as exc:
        wr.main(["--send", "--dry-run"])
    assert exc.value.code == 2
    assert wr.main(["--scheduled"]) == 0


def test_db_readonly_and_snapshot_isolation(tmp_path, inputs):
    path = tmp_path / "data.db"
    with sqlite3.connect(path) as conn:
        conn.execute("CREATE TABLE chat_store (create_time TEXT, in_robot INTEGER)")
        conn.execute("INSERT INTO chat_store VALUES (?, 1)", (WEEK,))
        conn.execute("CREATE TABLE timeout_event (created_at TEXT, theater TEXT, exclude_reason TEXT)")
        conn.execute("INSERT INTO timeout_event VALUES (?, '华东', '')", (WEEK,))
    before = path.read_bytes()
    args = {**inputs, "chat_store_json": None, "db_path": path}
    result = wr.run_weekly_report(WEEK, dry_run=True, **args)
    assert result["report"]["timeout"]["valid"] == 1
    assert path.read_bytes() == before
    result = wr.run_weekly_report(WEEK, dry_run=True, **inputs)
    assert result["report"]["timeout"] is None


def test_routes_require_session_and_prevent_traversal(tmp_path):
    app = Flask(__name__, template_folder=str(wr.ROOT / "templates"))
    app.config["WEEKLY_REPORT_DIR"] = str(tmp_path)
    register_weekly_reports(app)
    logged_in = False

    @app.before_request
    def login():
        if logged_in:
            g.mhero_session = SimpleNamespace(role="user")

    client = app.test_client()
    assert client.get("/weekly-reports").status_code == 401
    assert client.get(f"/weekly-reports/{WEEK}_weekly.xlsx").status_code == 401

    logged_in = True
    filename = f"{WEEK}_weekly.xlsx"
    (tmp_path / filename).write_bytes(b"test")
    assert client.get("/weekly-reports").status_code == 200
    response = client.get("/weekly-reports/" + filename)
    assert response.data == b"test" and response.headers["Cache-Control"] == "private, no-store"
    for name in ("../state.json", "%2e%2e/state.json", "state.json", "sub/" + filename):
        assert client.get("/weekly-reports/" + name).status_code == 404
    (tmp_path / filename).unlink()
    (tmp_path / filename).symlink_to(tmp_path.parent / "outside.xlsx")
    assert client.get("/weekly-reports/" + filename).status_code == 404


def test_real_app_route_registration():
    from app import app
    assert {"weekly_reports", "weekly_report_download"} <= {rule.endpoint for rule in app.url_map.iter_rules()}


def test_generation_failure_and_protected_state(inputs):
    inputs["metrics_jsonl"].write_text("bad json\n")
    with pytest.raises(ValueError):
        wr.run_weekly_report(WEEK, dry_run=True, **inputs)
    assert state.load_state(inputs["state_path"])[WEEK]["status"] == "failed"
    previous = {"status": "sent", "at": WEEK}
    state.save_state(inputs["state_path"], {WEEK: previous})
    with pytest.raises(ValueError):
        wr.run_weekly_report(WEEK, dry_run=True, **inputs)
    assert state.load_state(inputs["state_path"])[WEEK] == previous


def test_lock_across_processes(tmp_path):
    import subprocess
    import sys

    path = tmp_path / "state.json"
    with state.report_lock(path):
        child = subprocess.run([
            sys.executable, "-c",
            "from weekly_report_state import report_lock; import sys; "
            "\nwith report_lock(sys.argv[1]) as acquired: print(acquired)", str(path),
        ], cwd=wr.ROOT, check=True, capture_output=True, text=True)
        assert child.stdout.strip() == "False"
    with state.report_lock(path) as acquired:
        assert acquired


def test_feishu_client_uses_mocked_transport(monkeypatch):
    import io
    import urllib.error
    import feishu_client as fc

    monkeypatch.setenv("FEISHU_APP_ID", "test-app")
    monkeypatch.setenv("FEISHU_APP_SECRET", "test-secret")
    calls = []

    def fake_open(request, timeout):
        calls.append(request)
        payload = ({"code": 0, "tenant_access_token": "mock-token"} if len(calls) == 1
                   else {"code": 0, "data": {"message_id": "mock-id"}})
        return io.BytesIO(json.dumps(payload).encode())

    monkeypatch.setattr(fc.urllib.request, "urlopen", fake_open)
    assert fc.send_chat_text("mock-chat", "测试") == {"message_id": "mock-id"}
    assert len(calls) == 2
    assert json.loads(calls[0].data)["app_secret"] == "test-secret"
    assert calls[1].get_header("Authorization") == "Bearer mock-token"
    assert json.loads(json.loads(calls[1].data)["content"])["text"] == "测试"

    def fail(*args, **kwargs):
        raise urllib.error.URLError("test-secret mock-token")

    monkeypatch.setattr(fc.urllib.request, "urlopen", fail)
    with pytest.raises(RuntimeError) as exc:
        fc.send_chat_text("mock-chat", "测试")
    assert "test-secret" not in str(exc.value) and "mock-token" not in str(exc.value)
    with pytest.raises(ValueError):
        fc.send_chat_text("", "测试")
