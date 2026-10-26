"""拉群周报：纯函数统计、本地 Excel/简报与 CLI（默认 dry-run）。"""

from __future__ import annotations

import argparse
import json
import os
import re
import sqlite3
from datetime import date, datetime, time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font

from feishu_client import send_chat_text
from weekly_report_state import atomic_write, load_state, report_lock, save_state

ROOT = Path(__file__).resolve().parent
REPORT_DIR = ROOT / "data" / "weekly_reports"
REPORT_FILENAME_RE = re.compile(r"\d{4}-\d{2}-\d{2}_weekly\.xlsx")
SHANGHAI = ZoneInfo("Asia/Shanghai")
METRIC_KEYS = ("new_groups", "robot_rooms_added", "renamed_groups", "updated_groups", "in_robot_updated")


def parse_time(value):
    try:
        parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
    except (ValueError, TypeError):
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=SHANGHAI)
    return parsed.astimezone(SHANGHAI)


def week_bounds(week):
    monday = date.fromisoformat(str(week))
    if monday.weekday() != 0:
        raise ValueError("--week 必须是周一 YYYY-MM-DD")
    start = datetime.combine(monday, time.min, SHANGHAI)
    return start, start + timedelta(days=7)


def previous_week_start(today=None):
    today = today or datetime.now(SHANGHAI).date()
    return today - timedelta(days=today.weekday() + 7)


def load_config(path=None):
    cfg = json.loads((ROOT / "weekly_report.json").read_text(encoding="utf-8"))
    if path is not None:
        extra = json.loads(Path(path).read_text(encoding="utf-8"))
        if not isinstance(extra, dict):
            raise ValueError("周报配置必须为对象")
        cfg.update(extra)
    cfg["metrics_path"] = os.environ.get("STC_METRICS_PATH") or cfg["metrics_path"] or "~/m-hero-run/cron-logs/scrm-tier-metrics.jsonl"
    cfg["link_base"] = os.environ.get("STC_LINK_BASE") or cfg["link_base"]
    for key in ("enabled", "dry_run"):
        if type(cfg[key]) is not bool:
            raise ValueError(f"{key} 必须是布尔值")
    for key, upper in (("weekday", 7), ("hour", 24), ("minute", 60)):
        if type(cfg[key]) is not int or not 0 <= cfg[key] < upper:
            raise ValueError(f"非法周报时间：{key}")
    current = cfg["hour"] * 60 + cfg["minute"]
    for center in (0, 510, 540, 570, 720):
        diff = abs(current - center)
        if min(diff, 1440 - diff) <= 5:
            raise ValueError("周报时间必须避开 00:00/08:30/09:00/09:30/12:00 ±5 分钟")
    if type(cfg["sending_timeout_seconds"]) not in (int, float) or not 0 < cfg["sending_timeout_seconds"] < float("inf"):
        raise ValueError("sending_timeout_seconds 必须是有限正数")
    for key in ("chat_id", "link_base", "metrics_path", "state_path"):
        if not isinstance(cfg[key], str):
            raise ValueError(f"{key} 必须是字符串")
    if not cfg["state_path"].strip() or not cfg["chat_id"].strip():
        raise ValueError("state_path / chat_id 不能为空")
    return cfg


def _robot(row):
    return row.get("in_robot") in (1, "1", True)


def _coverage(rows):
    total = len(rows)
    robots = sum(_robot(row) for row in rows)
    return {
        "groups": total, "robot_groups": robots, "coverage_rate": robots / total if total else 0,
        "stores": len({str(row.get("store") or "").strip() for row in rows} - {""}),
        "robot_stores": len({str(row.get("store") or "").strip() for row in rows if _robot(row)} - {""}),
    }


def build_weekly_report(week_start, chat_store, metrics, timeout_events=None, timeout_summary=None):
    """不读磁盘/网络、不修改输入。半开周区间；覆盖为当前输入快照，非历史重建。"""
    start, stop = week_bounds(week_start)

    def within(value):
        at = parse_time(value)
        return at is not None and start <= at < stop

    new = [dict(row) for row in chat_store if within(row.get("create_time"))]
    new.sort(key=lambda row: (parse_time(row["create_time"]), str(row.get("room_id") or "")))
    theaters = []
    for name in sorted({str(row.get("theater") or "").strip() or "未归属战区" for row in chat_store}):
        rows = [row for row in chat_store if (str(row.get("theater") or "").strip() or "未归属战区") == name]
        weekly = [row for row in rows if within(row.get("create_time"))]
        theaters.append({"theater": name, **_coverage(rows), "new_groups": len(weekly),
                         "new_robot_groups": sum(_robot(row) for row in weekly)})
    runs = []
    for row in metrics:
        if row.get("tier") != "groups" or not within(row.get("start")):
            continue
        counts = row.get("counts") or {}
        runs.append({"start": parse_time(row["start"]).isoformat(), "status": row.get("status", ""),
                     **{key: int(counts.get(key, counts.get("groups_updated", 0) if key == "updated_groups" else 0) or 0)
                        for key in METRIC_KEYS}, "counts": dict(counts)})
    runs.sort(key=lambda row: row["start"])
    sync = {key: sum(row[key] for row in runs) for key in METRIC_KEYS}
    sync.update({"runs": len(runs), "failures": sum(row["status"] != "success" for row in runs),
                 "opening_counts": dict(runs[0]["counts"]) if runs else None,
                 "closing_counts": dict(runs[-1]["counts"]) if runs else None})
    timeout = None
    if timeout_summary is not None:
        event_counts = dict(timeout_summary["by_theater"])
        valid_counts = dict(timeout_summary["valid_by_theater"])
        timeout = {"events": int(timeout_summary["events"]), "valid": int(timeout_summary["valid"]),
                   "by_theater": [{"theater": name, "events": int(event_counts.get(name, 0)),
                                   "valid": int(valid_counts.get(name, 0))}
                                  for name in sorted(event_counts.keys() | valid_counts.keys())]}
        if (sum(row["events"] for row in timeout["by_theater"]) != timeout["events"]
                or sum(row["valid"] for row in timeout["by_theater"]) != timeout["valid"]
                or any(not 0 <= row["valid"] <= row["events"] for row in timeout["by_theater"])):
            raise ValueError("超时汇总总计与战区数据不一致")
    elif timeout_events is not None:
        counts = {}
        for row in timeout_events:
            if within(row.get("created_at")):
                name = str(row.get("theater") or "").strip() or "未归属战区"
                item = counts.setdefault(name, {"theater": name, "events": 0, "valid": 0})
                item["events"] += 1
                item["valid"] += not bool(row.get("exclude_reason"))
        timeout = {"events": sum(row["events"] for row in counts.values()),
                   "valid": sum(row["valid"] for row in counts.values()),
                   "by_theater": [counts[name] for name in sorted(counts)]}
    return {
        "week_start": start.date().isoformat(), "week_end": (stop - timedelta(days=1)).date().isoformat(),
        "coverage": _coverage(chat_store), "new_groups": new, "new_robot_groups": sum(_robot(row) for row in new),
        "theaters": theaters, "sync": sync, "sync_runs": runs, "timeout": timeout,
        "invalid_create_times": sum(parse_time(row.get("create_time")) is None for row in chat_store),
    }


def render_brief(report, download_url):
    c, s = report["coverage"], report["sync"]
    top = sorted((row for row in report["theaters"] if row["new_groups"]),
                 key=lambda row: (-row["new_groups"], row["theater"]))[:3]
    lines = [
        f"拉群覆盖周报｜{report['week_start']} ~ {report['week_end']}（北京时间）",
        f"期末（运行时快照）：客户群 {c['groups']}，已拉机器人 {c['robot_groups']}（{c['coverage_rate']:.1%}）；门店覆盖 {c['robot_stores']}/{c['stores']}",
        f"本周新建群 {len(report['new_groups'])}（其中当前已拉机器人 {report['new_robot_groups']}）；同步累计 robot_rooms_added {s['robot_rooms_added']}；同步失败 {s['failures']} 次",
        "口径：新建已拉机器人按创建时间及当前 in_robot；robot_rooms_added 是同步增量，含历史群，两者不同。",
        "战区新建 Top：" + ("；".join(f"{row['theater']} {row['new_groups']}" for row in top) or "无"),
    ]
    if s["runs"]:
        lines.append(f"groups 同步首末快照：{s['opening_counts'].get('groups', '缺失')} → {s['closing_counts'].get('groups', '缺失')}（周内首/末次记录）")
    else:
        lines.append("本周无 groups 同步记录（增量指标无数据，不代表实际为零）")
    t = report["timeout"]
    lines.append(f"本周超时事件 {t['events']}（有效 {t['valid']}）" if t is not None else "本周超时事件：未提供数据")
    if report["invalid_create_times"]:
        lines.append(f"创建时间缺失或无效 {report['invalid_create_times']} 群，未计入本周新建。")
    lines.append(f"明细下载（需登录）：{download_url}")
    return "\n".join(lines)


def write_weekly_excel(report, path):
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, headers, rows):
        ws = workbook.create_sheet(name)
        for values in [headers, *rows]:
            ws.append(values)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # 外部文本绝不作为公式执行。
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 24
        return ws

    sheet("本周新建群", ["room_id", "群名", "门店", "战区", "是否已拉机器人", "创建时间", "群主", "人数"], [
        [row.get("room_id", ""), row.get("group_name", ""), row.get("store", ""), row.get("theater", ""),
         "是" if _robot(row) else "否", row.get("create_time", ""), row.get("owner_name", ""), row.get("member_count", "")]
        for row in report["new_groups"]])
    coverage = sheet("战区覆盖", ["战区", "总群", "已拉机器人群", "覆盖率", "本周新建", "本周新建已拉机器人", "有群门店", "已拉机器人门店"], [
        [row[key] for key in ("theater", "groups", "robot_groups", "coverage_rate", "new_groups", "new_robot_groups", "stores", "robot_stores")]
        for row in report["theaters"]])
    for row in coverage.iter_rows(min_row=2):
        row[3].number_format = "0.0%"
    sheet("同步运行", ["时间", "status", *METRIC_KEYS, "groups 快照", "失败次数"], [
        [row["start"], row["status"], *[row[key] for key in METRIC_KEYS], row["counts"].get("groups"), int(row["status"] != "success")]
        for row in report["sync_runs"]] + [
        ["汇总", "", *[report["sync"][key] for key in METRIC_KEYS], "", report["sync"]["failures"]]])
    sheet("超时事件", ["战区", "本周事件", "有效超时"], [
        [row["theater"], row["events"], row["valid"]] for row in (report["timeout"] or {}).get("by_theater", [])])
    sheet("口径说明", ["项目", "说明"], [
        ["周期（北京时间）", f"{report['week_start']} 00:00:00 ~ {report['week_end']} 23:59:59"],
        ["覆盖快照", "当前 chat_store，不是历史周日快照；门店按非空店名去重"],
        ["机器人指标", "本周新建已拉机器人按 create_time + 当前 in_robot；robot_rooms_added 是同步增量，包含历史群"],
        ["同步首记录 counts", json.dumps(report["sync"]["opening_counts"], ensure_ascii=False)],
        ["同步末记录 counts", json.dumps(report["sync"]["closing_counts"], ensure_ascii=False)],
        ["来源", json.dumps(report.get("sources", {}), ensure_ascii=False)],
    ])
    try:
        atomic_write(path, workbook.save)
    finally:
        workbook.close()
    return Path(path)


def _load_json(path):
    return json.loads(Path(path).expanduser().read_text(encoding="utf-8"))


def _load_inputs(db_path, chat_store_json, metrics_path, timeout_summary_json):
    events = None
    rows = _load_json(chat_store_json) if chat_store_json else None
    # 显式快照模式不回读默认本地 DB，避免把两个时点的数据混在一起。
    if rows is None or db_path is not None:
        path = Path(db_path or os.environ.get("CLEANER_DB_PATH") or ROOT / "timeout.db").expanduser().resolve()
        with sqlite3.connect(path.as_uri() + "?mode=ro", uri=True) as conn:
            conn.row_factory = sqlite3.Row
            if rows is None:
                rows = [dict(row) for row in conn.execute("SELECT * FROM chat_store")]
            if not timeout_summary_json and conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='timeout_event'").fetchone():
                events = [dict(row) for row in conn.execute("SELECT created_at, theater, exclude_reason FROM timeout_event")]
    if isinstance(rows, dict):
        rows = rows.get("chat_store")
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError("chat_store 快照必须是对象列表或包含 chat_store 列表的对象")
    # JSONL 损坏则拒绝生成，不能悄悄少算失败/增量。
    metrics = []
    with Path(metrics_path).expanduser().open(encoding="utf-8") as fh:
        for line in fh:
            if line.strip():
                row = json.loads(line)
                if not isinstance(row, dict):
                    raise ValueError("metrics 每行必须为对象")
                if row.get("tier") == "groups" and parse_time(row.get("start")) is None:
                    raise ValueError("groups metrics 缺少有效 start")
                metrics.append(row)
    return rows, metrics, events, _load_json(timeout_summary_json) if timeout_summary_json else None


def run_weekly_report(week_start=None, *, dry_run=None, force=False, state_path=None,
                      chat_store_json=None, metrics_jsonl=None, timeout_summary_json=None,
                      db_path=None, out_dir=None, config_path=None):
    cfg = load_config(config_path)
    dry_run = cfg["dry_run"] if dry_run is None else dry_run
    if type(dry_run) is not bool or type(force) is not bool:
        raise ValueError("dry_run / force 必须为布尔值")
    week = week_bounds(week_start or previous_week_start())[0].date().isoformat()
    state_file = Path(state_path or cfg["state_path"]).expanduser()
    if not state_file.is_absolute():
        state_file = ROOT / state_file
    state_file = state_file.resolve()
    result = {"week_start": week, "dry_run": dry_run, "skipped": False, "state_path": str(state_file)}

    def skipped(reason):
        return {**result, "skipped": True, "skip_reason": reason}

    with report_lock(state_file) as acquired:
        if not acquired:
            return skipped("locked")
        state = load_state(state_file)
        previous = state.get(week)
        protected = previous is not None and previous["status"] in ("sent", "sending")
        if not dry_run and not force and previous:
            if previous["status"] == "sent":
                return skipped("sent")
            if previous["status"] == "sending":
                at = parse_time(previous.get("at"))
                if at is None or (datetime.now(SHANGHAI) - at).total_seconds() < cfg["sending_timeout_seconds"]:
                    return skipped("sending")

        def record(status, **extra):
            state[week] = {"status": status, "at": datetime.now(SHANGHAI).isoformat(), "chat_id": cfg["chat_id"], **extra}
            save_state(state_file, state)

        try:
            metrics_path = metrics_jsonl or cfg["metrics_path"]
            rows, metrics, events, summary = _load_inputs(db_path, chat_store_json, metrics_path, timeout_summary_json)
            report = build_weekly_report(week, rows, metrics, events, summary)
            report["sources"] = {"chat_store": str(chat_store_json or db_path or os.environ.get("CLEANER_DB_PATH") or ROOT / "timeout.db"),
                                 "metrics": str(metrics_path), "timeout_summary": str(timeout_summary_json or ""),
                                 "generated_at": datetime.now(SHANGHAI).isoformat()}
            directory = Path(out_dir) if out_dir is not None else REPORT_DIR
            filename = f"{week}_weekly.xlsx"
            excel = write_weekly_excel(report, directory / filename)
            brief = render_brief(report, f"{cfg['link_base'].rstrip('/')}/weekly-reports/{filename}")
            brief_path = directory / "brief.txt"
            atomic_write(brief_path, lambda tmp: tmp.write_text(brief + "\n", encoding="utf-8"))
            atomic_write(directory / f"{week}_brief.txt", lambda tmp: tmp.write_text(brief + "\n", encoding="utf-8"))
            result.update(report=report, brief=brief, excel_path=str(excel), brief_path=str(brief_path))
        except Exception:
            if not protected:
                record("failed")
            raise
        if dry_run:
            if not protected:
                record("dry_run")
            return result
        record("sending")  # 必须先持久化；最终 sent 写失败时保留 sending。
        try:
            message = send_chat_text(cfg["chat_id"], brief)
        except Exception:
            if protected:
                state[week] = previous
                save_state(state_file, state)
            else:
                record("failed")
            raise
        record("sent", message_id=(message or {}).get("message_id"))
        return result


def main(argv=None):
    parser = argparse.ArgumentParser(description="拉群覆盖上一完整自然周周报，默认 dry-run")
    parser.add_argument("--week")
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument("--dry-run", dest="dry_run", action="store_const", const=True, default=None)
    modes.add_argument("--send", dest="dry_run", action="store_const", const=False)
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--scheduled", action="store_true", help="仅 enabled=true 且到配置的星期/时间才运行，供每日 cron")
    for flag in ("config", "state-path", "chat-store-json", "metrics-jsonl", "timeout-summary-json", "db-path", "out"):
        parser.add_argument("--" + flag, type=Path)
    args = parser.parse_args(argv)
    try:
        cfg = load_config(args.config)
        now = datetime.now(SHANGHAI)
        if args.scheduled and (not cfg["enabled"] or (now.weekday(), now.hour, now.minute) != (cfg["weekday"], cfg["hour"], cfg["minute"])):
            print("周报调度未启用或未到计划时间，跳过")
            return 0
        result = run_weekly_report(args.week, dry_run=args.dry_run, force=args.force, config_path=args.config,
                                   state_path=args.state_path, chat_store_json=args.chat_store_json,
                                   metrics_jsonl=args.metrics_jsonl, timeout_summary_json=args.timeout_summary_json,
                                   db_path=args.db_path, out_dir=args.out)
    except (ValueError, OSError, RuntimeError, sqlite3.Error) as exc:
        parser.exit(1, f"周报失败：{exc}\n")
    if result["skipped"]:
        print(f"周报跳过：{result['skip_reason']}")
    else:
        print("模式：" + ("dry-run" if result["dry_run"] else "已发送"))
        print(result["brief"])
        print(f"Excel：{result['excel_path']}\n简报：{result['brief_path']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
